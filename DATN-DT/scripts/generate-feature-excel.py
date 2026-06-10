# -*- coding: utf-8 -*-
"""Generate Excel listing all features, files, and flows for AloBooking project."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = os.path.join(os.path.dirname(__file__), '..', 'AloBooking_DanhSachChucNang.xlsx')

HEADER_FILL = PatternFill('solid', fgColor='16A34A')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
THIN = Side(style='thin')


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def auto_width(ws, min_w=12, max_w=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letter].width = min(max(length + 2, min_w), max_w)


def add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    style_header(ws)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'
    auto_width(ws)
    return ws


# ── Sheet 1: Tổng quan module ──
modules = [
    ('1', 'Xác thực & Tài khoản', 'Đăng ký, đăng nhập, Google OAuth, quên MK, profile, admin user', 'Guest, User, Admin', 'users'),
    ('2', 'Sân bóng (Fields)', 'Danh sách, chi tiết, CRUD sân, upload ảnh Cloudinary', 'Guest, User, Admin', 'fields'),
    ('3', 'Bảng giá khung giờ', 'CRUD giá theo thứ + khung giờ', 'Guest, Admin', 'field-prices'),
    ('4', 'Đặt sân & Thanh toán', 'Chọn slot, giữ realtime, checkout, cash/MoMo/VNPay, hủy đơn', 'User, Admin', 'bookings'),
    ('5', 'Mã giảm giá', 'Validate mã, CRUD khuyến mãi', 'User, Admin', 'discounts'),
    ('6', 'Đánh giá (Review)', 'Xem/tạo/xóa review, phản hồi admin', 'Guest, User', 'reviews'),
    ('7', 'Blog / Tin tức', 'Danh sách, chi tiết, CRUD admin', 'Guest, Admin', 'blog'),
    ('8', 'Thông báo', 'In-app notification, realtime socket, admin broadcast', 'User, Admin', 'notifications'),
    ('9', 'Chatbot AI', 'AloBookingBot: đặt sân, tra cứu, hủy, Groq AI', 'User', 'chatbot'),
    ('10', 'Dashboard Admin', 'KPI, biểu đồ doanh thu, top sân', 'Admin', 'dashboard'),
    ('11', 'Socket.IO Realtime', 'Giữ slot, push notification', 'User', 'socket'),
    ('12', 'Cron Job', 'Nhắc lịch trước giờ đá (30 phút/lần)', 'System', 'cron'),
    ('13', 'Trang tĩnh / Marketing', 'Trang chủ, FAQ, hướng dẫn, chính sách', 'Guest', 'static'),
    ('14', 'Hạ tầng', 'Server, CORS, DB, auth, deploy', 'System', 'infra'),
]

# ── Sheet 2: File Backend ──
backend_files = [
    ('Xác thực', 'Route', 'backend/src/routes/users.routes.js', 'Định nghĩa endpoint /api/users/*'),
    ('Xác thực', 'Controller', 'backend/src/controller/user.controller.js', 'Xử lý request auth, profile, chatbot'),
    ('Xác thực', 'Service', 'backend/src/services/users.service.js', 'Logic đăng ký, login, OTP, admin user'),
    ('Xác thực', 'Model', 'backend/src/models/users.model.js', 'Schema User'),
    ('Xác thực', 'Model', 'backend/src/models/apiKey.model.js', 'Phiên đăng nhập / refresh token'),
    ('Xác thực', 'Model', 'backend/src/models/otp.model.js', 'OTP quên mật khẩu'),
    ('Xác thực', 'Middleware', 'backend/src/auth/checkAuth.js', 'authUser, authAdmin, asyncHandler'),
    ('Xác thực', 'Utils', 'backend/src/utils/jwt.js', 'Tạo/verify JWT'),
    ('Xác thực', 'Utils', 'backend/src/utils/getAuthToken.js', 'Đọc token cookie/header'),
    ('Xác thực', 'Utils', 'backend/src/utils/cookieOptions.js', 'SameSite cookie cross-origin'),
    ('Xác thực', 'Utils', 'backend/src/utils/sendMailForgotPassword.js', 'Gửi email OTP'),
    ('Xác thực', 'Utils', 'backend/src/utils/normalizeEmail.js', 'Chuẩn hóa email'),
    ('Xác thực', 'Config', 'backend/src/config/superAdmin.js', 'Email Super Admin cố định'),
    ('Sân bóng', 'Route', 'backend/src/routes/field.routes.js', 'CRUD /api/fields/*'),
    ('Sân bóng', 'Controller', 'backend/src/controller/field.controller.js', ''),
    ('Sân bóng', 'Service', 'backend/src/services/field.service.js', 'Upload/xóa ảnh Cloudinary'),
    ('Sân bóng', 'Model', 'backend/src/models/field.model.js', 'Schema sân 5/7/11'),
    ('Sân bóng', 'Config', 'backend/src/config/cloudinaryUpload.js', 'Upload ảnh'),
    ('Sân bóng', 'Config', 'backend/src/config/cloudDinary.js', 'Cấu hình Cloudinary'),
    ('Bảng giá', 'Route', 'backend/src/routes/fieldPrice.routes.js', ''),
    ('Bảng giá', 'Controller', 'backend/src/controller/fieldPrice.controller.js', ''),
    ('Bảng giá', 'Service', 'backend/src/services/fieldPrice.service.js', ''),
    ('Bảng giá', 'Model', 'backend/src/models/fieldPrice.model.js', ''),
    ('Đặt sân', 'Route', 'backend/src/routes/booking.routes.js', 'create, cancel, verify payment, admin'),
    ('Đặt sân', 'Controller', 'backend/src/controller/booking.controller.js', ''),
    ('Đặt sân', 'Service', 'backend/src/services/booking.service.js', 'Tạo đơn, MoMo/VNPay, hủy'),
    ('Đặt sân', 'Model', 'backend/src/models/booking.model.js', ''),
    ('Mã giảm giá', 'Route', 'backend/src/routes/discount.routes.js', ''),
    ('Mã giảm giá', 'Controller', 'backend/src/controller/discount.controller.js', ''),
    ('Mã giảm giá', 'Service', 'backend/src/services/discount.service.js', ''),
    ('Mã giảm giá', 'Model', 'backend/src/models/discount.model.js', ''),
    ('Đánh giá', 'Route', 'backend/src/routes/review.routes.js', ''),
    ('Đánh giá', 'Controller', 'backend/src/controller/review.controller.js', ''),
    ('Đánh giá', 'Service', 'backend/src/services/review.service.js', ''),
    ('Đánh giá', 'Model', 'backend/src/models/review.model.js', ''),
    ('Blog', 'Route', 'backend/src/routes/blog.routes.js', ''),
    ('Blog', 'Controller', 'backend/src/controller/blog.controller.js', ''),
    ('Blog', 'Service', 'backend/src/services/blog.service.js', ''),
    ('Blog', 'Model', 'backend/src/models/blog.model.js', ''),
    ('Thông báo', 'Route', 'backend/src/routes/notification.routes.js', ''),
    ('Thông báo', 'Controller', 'backend/src/controller/notification.controller.js', ''),
    ('Thông báo', 'Service', 'backend/src/services/notification.service.js', 'Tạo + socket push'),
    ('Thông báo', 'Model', 'backend/src/models/notification.model.js', ''),
    ('Chatbot', 'Service', 'backend/src/services/chatbot.service.js', 'processMessage, flow đặt/hủy'),
    ('Chatbot', 'Utils', 'backend/src/utils/chatbotIntents.js', 'Nhận diện intent'),
    ('Chatbot', 'Utils', 'backend/src/utils/chatbotData.js', 'Query DB cho bot'),
    ('Chatbot', 'Utils', 'backend/src/utils/chatbot.js', 'Groq AI polish'),
    ('Chatbot', 'Model', 'backend/src/models/chatbotSession.model.js', ''),
    ('Chatbot', 'Model', 'backend/src/models/messageChatbot.model.js', ''),
    ('Dashboard', 'Route', 'backend/src/routes/dashboard.routes.js', ''),
    ('Dashboard', 'Controller', 'backend/src/controller/dashboard.controller.js', ''),
    ('Dashboard', 'Service', 'backend/src/services/dashboard.service.js', ''),
    ('Socket', 'Config', 'backend/src/config/socket.js', 'hold-slots, notification push'),
    ('Socket', 'Service', 'backend/src/services/slotHolding.service.js', 'In-memory giữ slot 5 phút'),
    ('Cron', 'Job', 'backend/src/jobs/bookingReminder.js', 'Nhắc lịch mỗi 30 phút'),
    ('Hạ tầng', 'Entry', 'backend/src/server.js', 'Bootstrap Express + Socket + Cron'),
    ('Hạ tầng', 'Routes', 'backend/src/routes/index.routes.js', 'Gom tất cả /api/*'),
    ('Hạ tầng', 'Config', 'backend/src/config/connectDB.js', 'MongoDB'),
    ('Hạ tầng', 'Config', 'backend/src/config/cors.js', 'CORS origins'),
    ('Hạ tầng', 'Config', 'backend/src/config/serveFrontend.js', 'Serve SPA production'),
    ('Hạ tầng', 'Core', 'backend/src/core/success.response.js', ''),
    ('Hạ tầng', 'Core', 'backend/src/core/error.response.js', ''),
    ('Hạ tầng', 'Core', 'backend/src/core/statusCodes.js', ''),
    ('Hạ tầng', 'Seed', 'backend/src/seed/seedDatabase.js', 'npm run seed'),
    ('Hạ tầng', 'Seed', 'backend/src/seed/fixDuplicateEmails.js', 'Dọn email trùng'),
]

# ── Sheet 3: File Frontend ──
frontend_files = [
    ('Entry & Router', 'Entry', 'frontend/src/main.jsx', 'Khởi tạo React + Router'),
    ('Entry & Router', 'Routes', 'frontend/src/routes/index.jsx', 'Định nghĩa tất cả route'),
    ('Entry & Router', 'Guard', 'frontend/src/components/RouteGuard.jsx', 'AdminRoute, UserRoute, GuestRoute'),
    ('State', 'Store', 'frontend/src/store/Provider.jsx', 'fetchAuth, dataUser global'),
    ('State', 'Context', 'frontend/src/store/Context.jsx', ''),
    ('State', 'Hook', 'frontend/src/hooks/useStore.jsx', ''),
    ('HTTP', 'Client', 'frontend/src/config/axiosClient.jsx', 'Axios + refresh token'),
    ('HTTP', 'Base', 'frontend/src/config/request.jsx', 'ApiClient instance'),
    ('HTTP', 'Auth', 'frontend/src/utils/authToken.js', 'Bearer token localStorage'),
    ('Xác thực', 'Page', 'frontend/src/pages/LoginUser.jsx', 'Đăng nhập email + Google'),
    ('Xác thực', 'Page', 'frontend/src/pages/RegisterUser.jsx', 'Đăng ký'),
    ('Xác thực', 'Page', 'frontend/src/pages/ForgotPassword.jsx', 'OTP quên MK'),
    ('Xác thực', 'Page', 'frontend/src/pages/Profile.jsx', 'Hồ sơ, lịch sử, đổi MK'),
    ('Xác thực', 'API', 'frontend/src/config/UserRequest.jsx', 'Auth, profile, chatbot'),
    ('Xác thực', 'API', 'frontend/src/config/AdminUserRequest.js', 'Admin CRUD user'),
    ('Sân bóng', 'Page', 'frontend/src/pages/FieldList.jsx', 'Danh sách sân'),
    ('Sân bóng', 'Page', 'frontend/src/pages/DetailField.jsx', 'Chi tiết + chọn slot'),
    ('Sân bóng', 'Component', 'frontend/src/components/FeaturedFields.jsx', 'Sân nổi bật trang chủ'),
    ('Sân bóng', 'Component', 'frontend/src/components/FieldCard/index.jsx', 'Card sân'),
    ('Sân bóng', 'Component', 'frontend/src/components/Banner.jsx', 'Tìm sân trang chủ'),
    ('Sân bóng', 'API', 'frontend/src/config/FieldRequest.js', ''),
    ('Sân bóng', 'Admin', 'frontend/src/pages/admin/FieldManagement.jsx', 'CRUD sân'),
    ('Bảng giá', 'Admin', 'frontend/src/pages/admin/FieldPriceManagement.jsx', ''),
    ('Bảng giá', 'API', 'frontend/src/config/FieldPriceRequest.js', ''),
    ('Đặt sân', 'Page', 'frontend/src/pages/Checkout.jsx', 'Thanh toán + mã giảm giá'),
    ('Đặt sân', 'Page', 'frontend/src/pages/BookingSuccess.jsx', 'Kết quả đặt sân'),
    ('Đặt sân', 'API', 'frontend/src/config/BookingRequest.js', ''),
    ('Đặt sân', 'Socket', 'frontend/src/hooks/useSlotSocket.js', 'Giữ slot realtime'),
    ('Đặt sân', 'Admin', 'frontend/src/pages/admin/BookingManagement.jsx', ''),
    ('Mã giảm giá', 'API', 'frontend/src/config/DiscountRequest.js', ''),
    ('Mã giảm giá', 'Admin', 'frontend/src/pages/admin/DiscountManagement.jsx', ''),
    ('Đánh giá', 'Component', 'frontend/src/components/Feedback.jsx', 'Carousel review trang chủ'),
    ('Đánh giá', 'Component', 'frontend/src/components/ReviewModal.jsx', 'Modal đánh giá'),
    ('Đánh giá', 'API', 'frontend/src/config/ReviewRequest.js', ''),
    ('Blog', 'Page', 'frontend/src/pages/BlogPage.jsx', ''),
    ('Blog', 'Page', 'frontend/src/pages/BlogDetail.jsx', ''),
    ('Blog', 'Component', 'frontend/src/components/HomeBlog.jsx', '3 bài trang chủ'),
    ('Blog', 'API', 'frontend/src/config/BlogRequest.jsx', ''),
    ('Blog', 'Admin', 'frontend/src/pages/admin/BlogAdmin.jsx', ''),
    ('Thông báo', 'Page', 'frontend/src/pages/NotificationsPage.jsx', ''),
    ('Thông báo', 'Component', 'frontend/src/components/Header.jsx', 'Dropdown + badge'),
    ('Thông báo', 'API', 'frontend/src/config/NotificationRequest.js', ''),
    ('Thông báo', 'Socket', 'frontend/src/hooks/useSocket.js', 'useNotificationSocket'),
    ('Thông báo', 'Admin', 'frontend/src/pages/admin/NotificationAdmin.jsx', ''),
    ('Chatbot', 'Component', 'frontend/src/components/ChatBot.jsx', 'Widget nổi toàn site'),
    ('Chatbot', 'Layout', 'frontend/src/layouts/RootLayout.jsx', 'Gắn ChatBot global'),
    ('Dashboard', 'Page', 'frontend/src/pages/admin/Dashboard.jsx', ''),
    ('Dashboard', 'API', 'frontend/src/config/DashboardRequest.js', ''),
    ('Admin Layout', 'Layout', 'frontend/src/layouts/AdminLayout.jsx', ''),
    ('Admin Layout', 'Component', 'frontend/src/components/admin/Sidebar.jsx', ''),
    ('Admin Layout', 'Component', 'frontend/src/components/admin/AdminHeader.jsx', ''),
    ('Trang chủ', 'Page', 'frontend/src/App.jsx', 'Trang chủ'),
    ('Trang tĩnh', 'Page', 'frontend/src/pages/GuidePage.jsx', 'Hướng dẫn'),
    ('Trang tĩnh', 'Page', 'frontend/src/pages/FAQPage.jsx', 'FAQ'),
    ('Trang tĩnh', 'Page', 'frontend/src/pages/PrivacyPage.jsx', 'Chính sách'),
    ('Layout', 'Component', 'frontend/src/components/Footer.jsx', ''),
    ('Utils', 'Utils', 'frontend/src/utils/imageUrl.js', 'URL ảnh'),
    ('Utils', 'Hook', 'frontend/src/hooks/useDebounce.jsx', 'Debounce search'),
    ('Config', 'Config', 'frontend/src/config/superAdmin.js', 'Super Admin UI'),
]

# ── Sheet 4: API Endpoints ──
api_endpoints = [
    ('Users', 'POST', '/api/users/register', 'Công khai', 'Đăng ký'),
    ('Users', 'POST', '/api/users/login', 'Công khai', 'Đăng nhập'),
    ('Users', 'POST', '/api/users/login-google', 'Công khai', 'Google OAuth'),
    ('Users', 'GET', '/api/users/auth', 'User', 'Lấy thông tin phiên'),
    ('Users', 'GET', '/api/users/refresh-token', 'Công khai', 'Refresh JWT'),
    ('Users', 'POST', '/api/users/logout', 'User', 'Đăng xuất'),
    ('Users', 'POST', '/api/users/forgot-password', 'Công khai', 'Gửi OTP'),
    ('Users', 'POST', '/api/users/reset-password', 'Công khai', 'Đặt lại MK'),
    ('Users', 'PUT', '/api/users/update', 'User', 'Cập nhật profile'),
    ('Users', 'PUT', '/api/users/change-password', 'User', 'Đổi mật khẩu'),
    ('Users', 'POST', '/api/users/upload-avatar', 'User', 'Upload avatar'),
    ('Users', 'POST', '/api/users/chatbot', 'User', 'Gửi tin chatbot'),
    ('Users', 'GET', '/api/users/message-chatbot', 'User', 'Lịch sử chatbot'),
    ('Users', 'GET', '/api/users/admin/users', 'Admin', 'Danh sách user'),
    ('Users', 'PUT', '/api/users/admin/users/:id', 'Admin', 'Sửa user'),
    ('Users', 'DELETE', '/api/users/admin/users/:id', 'Admin', 'Xóa user'),
    ('Fields', 'GET', '/api/fields/', 'Công khai', 'Danh sách sân'),
    ('Fields', 'GET', '/api/fields/:id', 'Công khai', 'Chi tiết sân'),
    ('Fields', 'POST', '/api/fields/create', 'Admin', 'Tạo sân'),
    ('Fields', 'PUT', '/api/fields/update/:id', 'Admin', 'Sửa sân'),
    ('Fields', 'DELETE', '/api/fields/delete/:id', 'Admin', 'Xóa sân'),
    ('Fields', 'PATCH', '/api/fields/status/:id', 'Admin', 'Đổi trạng thái'),
    ('Field Prices', 'GET', '/api/field-prices/:fieldId', 'Công khai', 'Giá theo sân'),
    ('Field Prices', 'POST', '/api/field-prices/create', 'Công khai*', 'Tạo giá'),
    ('Field Prices', 'PUT', '/api/field-prices/update/:id', 'Admin', 'Sửa giá'),
    ('Field Prices', 'DELETE', '/api/field-prices/delete/:id', 'Admin', 'Xóa giá'),
    ('Bookings', 'POST', '/api/bookings/create', 'User', 'Tạo đơn đặt sân'),
    ('Bookings', 'GET', '/api/bookings/field/:fieldId', 'Công khai', 'Slot đã đặt theo ngày'),
    ('Bookings', 'GET', '/api/bookings/my-bookings', 'User', 'Lịch sử của tôi'),
    ('Bookings', 'GET', '/api/bookings/:id', 'Công khai', 'Chi tiết đơn'),
    ('Bookings', 'PUT', '/api/bookings/cancel/:id', 'User', 'Hủy đơn'),
    ('Bookings', 'PUT', '/api/bookings/verify-momo/:bookingId', 'User', 'Xác minh MoMo'),
    ('Bookings', 'PUT', '/api/bookings/verify-vnpay/:bookingId', 'User', 'Xác minh VNPay'),
    ('Bookings', 'GET', '/api/bookings/admin/all', 'Admin', 'Tất cả đơn'),
    ('Bookings', 'PUT', '/api/bookings/admin/status/:bookingId', 'Admin', 'Đổi trạng thái'),
    ('Bookings', 'PUT', '/api/bookings/admin/cancel/:id', 'Admin', 'Hủy đơn admin'),
    ('Discounts', 'POST', '/api/discounts/validate', 'User', 'Validate mã'),
    ('Discounts', 'GET', '/api/discounts/available', 'User', 'Mã khả dụng'),
    ('Discounts', 'GET', '/api/discounts/', 'Admin', 'Danh sách mã'),
    ('Discounts', 'POST', '/api/discounts/create', 'Admin', 'Tạo mã'),
    ('Discounts', 'PUT', '/api/discounts/:id', 'Admin', 'Sửa mã'),
    ('Discounts', 'DELETE', '/api/discounts/:id', 'Admin', 'Xóa mã'),
    ('Reviews', 'GET', '/api/reviews/all', 'Công khai', 'Tất cả review'),
    ('Reviews', 'GET', '/api/reviews/field/:fieldId', 'Công khai', 'Review theo sân'),
    ('Reviews', 'POST', '/api/reviews/create', 'User', 'Tạo review'),
    ('Reviews', 'GET', '/api/reviews/can-review/:bookingId', 'User', 'Kiểm tra quyền'),
    ('Reviews', 'DELETE', '/api/reviews/:reviewId', 'User', 'Xóa review'),
    ('Reviews', 'POST', '/api/reviews/:reviewId/reply', 'User', 'Phản hồi review'),
    ('Blog', 'GET', '/api/blog/get-all', 'Công khai', 'Danh sách blog'),
    ('Blog', 'GET', '/api/blog/get-by-id', 'Công khai', 'Chi tiết blog'),
    ('Blog', 'POST', '/api/blog/create', 'Admin', 'Tạo bài'),
    ('Blog', 'POST', '/api/blog/update/:id', 'Admin', 'Sửa bài'),
    ('Blog', 'DELETE', '/api/blog/delete/:id', 'Admin', 'Xóa bài'),
    ('Blog', 'POST', '/api/blog/upload-image', 'Admin', 'Upload ảnh blog'),
    ('Notifications', 'GET', '/api/notifications/all', 'User', 'Danh sách'),
    ('Notifications', 'GET', '/api/notifications/unread-count', 'User', 'Số chưa đọc'),
    ('Notifications', 'POST', '/api/notifications/:id/read', 'User', 'Đánh dấu đọc'),
    ('Notifications', 'POST', '/api/notifications/read-all', 'User', 'Đọc tất cả'),
    ('Notifications', 'DELETE', '/api/notifications/:id', 'User', 'Xóa'),
    ('Notifications', 'POST', '/api/notifications/admin/promotion', 'Admin', 'Gửi KM'),
    ('Notifications', 'POST', '/api/notifications/admin/system', 'Admin', 'Gửi hệ thống'),
    ('Dashboard', 'GET', '/api/dashboard/stats', 'Admin', 'KPI'),
    ('Dashboard', 'GET', '/api/dashboard/revenue-chart', 'Admin', 'Biểu đồ doanh thu'),
    ('Dashboard', 'GET', '/api/dashboard/field-distribution', 'Admin', 'Phân bố loại sân'),
    ('Dashboard', 'GET', '/api/dashboard/top-fields', 'Admin', 'Top sân'),
    ('Dashboard', 'GET', '/api/dashboard/recent-bookings', 'Admin', 'Đơn gần đây'),
    ('Dashboard', 'GET', '/api/dashboard/full', 'Admin', 'Dashboard đầy đủ'),
    ('System', 'GET', '/health', 'Công khai', 'Health check'),
]

# ── Sheet 5: Luồng chạy ──
flows = [
    ('Đăng nhập', '1', 'LoginUser.jsx', 'UserRequest.login()', 'POST /api/users/login'),
    ('Đăng nhập', '2', 'user.controller.login', 'users.service.login', 'Set cookie JWT + metadata AES'),
    ('Đăng nhập', '3', 'Provider.fetchAuth()', 'Giải mã AES → dataUser', 'RouteGuard → / hoặc /admin'),
    ('Đăng ký', '1', 'RegisterUser.jsx', 'POST /api/users/register', 'users.service.createUser + notification chào mừng'),
    ('Quên MK', '1', 'ForgotPassword.jsx', 'POST /api/users/forgot-password', 'Gửi OTP email'),
    ('Quên MK', '2', 'ForgotPassword.jsx', 'POST /api/users/reset-password', 'OTP + mật khẩu mới'),
    ('Xem sân', '1', 'FieldList.jsx / DetailField.jsx', 'GET /api/fields/', 'field.service'),
    ('Xem sân', '2', 'DetailField.jsx', 'GET /api/field-prices/:fieldId', 'Hiển thị bảng giá'),
    ('Xem sân', '3', 'DetailField.jsx', 'GET /api/bookings/field/:fieldId?date=', 'Slot đã đặt'),
    ('Giữ slot', '1', 'DetailField.jsx → useSlotSocket.js', 'socket.emit(hold-slots)', 'slotHolding.service (5 phút)'),
    ('Giữ slot', '2', 'config/socket.js', 'broadcast slots-held-by-others', 'Cập nhật UI realtime'),
    ('Checkout', '1', 'Checkout.jsx', 'POST /api/discounts/validate', 'Áp mã giảm giá'),
    ('Checkout', '2', 'Checkout.jsx', 'POST /api/bookings/create', 'booking.service.createBooking'),
    ('Checkout', '3a', 'Tiền mặt', 'status pending', '→ /booking-success/:id'),
    ('Checkout', '3b', 'MoMo/VNPay', 'Redirect cổng thanh toán', 'PUT verify-momo|vnpay → paid'),
    ('Checkout', '4', 'socket booking-confirmed', 'notification.service', 'Push thông báo xác nhận'),
    ('Hủy đơn', '1', 'Profile.jsx', 'PUT /api/bookings/cancel/:id', 'booking.service + notification'),
    ('Chatbot', '1', 'ChatBot.jsx', 'POST /api/users/chatbot', 'chatbot.service.processMessage'),
    ('Chatbot', '2', 'chatbotIntents.js', 'Nhận diện intent', 'booking/availability/price/...'),
    ('Chatbot', '3', 'chatbotData.js', 'Query DB', 'Field, Booking, Discount'),
    ('Chatbot', '4', 'chatbot.js', 'Groq AI polish', 'Lưu MessageChatbot'),
    ('Thông báo RT', '1', 'notification.service', 'socket emit new-notification', 'Header useNotificationSocket'),
    ('Cron nhắc lịch', '1', 'bookingReminder.js', 'Mỗi 30 phút', 'Booking hôm nay 30-90p trước giờ đá'),
    ('Cron nhắc lịch', '2', 'notification.service', 'booking_reminder', 'reminderSent = true'),
    ('Admin Dashboard', '1', 'Dashboard.jsx', 'GET /api/dashboard/full', 'dashboard.service aggregate'),
    ('Admin CRUD sân', '1', 'FieldManagement.jsx', 'POST/PUT/DELETE /api/fields/*', 'Cloudinary upload'),
    ('Đánh giá', '1', 'Profile.jsx → ReviewModal.jsx', 'POST /api/reviews/create', 'Cập nhật rating Field'),
    ('Khởi động server', '1', 'server.js', 'bootstrap()', 'routes → serveFrontend → socket → cron'),
    ('Khởi động FE', '1', 'main.jsx', 'Provider.fetchAuth()', 'RouterProvider → RootLayout → Page'),
]

# ── Sheet 6: Route Frontend ──
fe_routes = [
    ('/', 'App.jsx', 'Công khai', 'Trang chủ'),
    ('/fields', 'FieldList.jsx', 'Công khai', 'Danh sách sân'),
    ('/san/:id', 'DetailField.jsx', 'Công khai', 'Chi tiết sân + đặt'),
    ('/blogs', 'BlogPage.jsx', 'Công khai', 'Tin tức'),
    ('/blog/:id', 'BlogDetail.jsx', 'Công khai', 'Chi tiết blog'),
    ('/huong-dan', 'GuidePage.jsx', 'Công khai', 'Hướng dẫn'),
    ('/cau-hoi-thuong-gap', 'FAQPage.jsx', 'Công khai', 'FAQ'),
    ('/policy', 'PrivacyPage.jsx', 'Công khai', 'Chính sách'),
    ('/login', 'LoginUser.jsx', 'Guest', 'Đăng nhập'),
    ('/register', 'RegisterUser.jsx', 'Guest', 'Đăng ký'),
    ('/forgot-password', 'ForgotPassword.jsx', 'Guest', 'Quên MK'),
    ('/checkout', 'Checkout.jsx', 'User', 'Thanh toán'),
    ('/booking-success/:id', 'BookingSuccess.jsx', 'User', 'Kết quả đặt sân'),
    ('/profile', 'Profile.jsx', 'User', 'Tài khoản'),
    ('/notifications', 'NotificationsPage.jsx', 'User', 'Thông báo'),
    ('/admin', 'Dashboard.jsx', 'Admin', 'Thống kê'),
    ('/admin/fields', 'FieldManagement.jsx', 'Admin', 'Quản lý sân'),
    ('/admin/field-prices', 'FieldPriceManagement.jsx', 'Admin', 'Cấu hình giá'),
    ('/admin/bookings', 'BookingManagement.jsx', 'Admin', 'Quản lý đơn'),
    ('/admin/discounts', 'DiscountManagement.jsx', 'Admin', 'Mã giảm giá'),
    ('/admin/users', 'UserManagement.jsx', 'Admin', 'Người dùng'),
    ('/admin/blogs', 'BlogAdmin.jsx', 'Admin', 'Bài viết'),
    ('/admin/notifications', 'NotificationAdmin.jsx', 'Admin', 'Gửi thông báo'),
]

# ── Sheet 7: Socket Events ──
socket_events = [
    ('Client → Server', 'join-user', 'Vào room user_{userId}', 'useSocket.js / Header.jsx'),
    ('Client → Server', 'join-field', 'Vào room field_{fieldId}_{date}', 'useSlotSocket.js'),
    ('Client → Server', 'leave-field', 'Rời room sân', 'useSlotSocket.js'),
    ('Client → Server', 'hold-slots', 'Giữ tạm slot (5 phút)', 'DetailField.jsx'),
    ('Client → Server', 'release-slots', 'Bỏ giữ slot', 'DetailField.jsx'),
    ('Client → Server', 'booking-confirmed', 'Xác nhận đã đặt', 'Checkout.jsx'),
    ('Client → Server', 'disconnect', 'Tự giải phóng slot', 'socket.js'),
    ('Server → Client', 'held-slots', 'Danh sách slot đang giữ', 'useSlotSocket.js'),
    ('Server → Client', 'hold-result', 'Kết quả giữ slot', 'useSlotSocket.js'),
    ('Server → Client', 'slots-held-by-others', 'Slot bị người khác giữ', 'useSlotSocket.js'),
    ('Server → Client', 'slots-released', 'Slot được giải phóng', 'useSlotSocket.js'),
    ('Server → Client', 'slots-booked', 'Slot đã được đặt', 'useSlotSocket.js'),
    ('Server → Client', 'new-notification', 'Thông báo mới', 'useNotificationSocket'),
]

# ── Build workbook ──
wb = Workbook()
wb.remove(wb.active)

add_sheet(wb, 'Tổng quan Module', ['STT', 'Module', 'Mô tả', 'Vai trò', 'API prefix'], modules)
add_sheet(wb, 'Backend Files', ['Module', 'Loại file', 'Đường dẫn', 'Ghi chú'], backend_files)
add_sheet(wb, 'Frontend Files', ['Module', 'Loại file', 'Đường dẫn', 'Ghi chú'], frontend_files)
add_sheet(wb, 'API Endpoints', ['Module', 'Method', 'Endpoint', 'Quyền', 'Mô tả'], api_endpoints)
add_sheet(wb, 'Luồng chạy', ['Chức năng', 'Bước', 'File/Component', 'Hành động tiếp', 'Kết quả'], flows)
add_sheet(wb, 'Frontend Routes', ['Route', 'Component', 'Quyền', 'Mô tả'], fe_routes)
add_sheet(wb, 'Socket Events', ['Hướng', 'Event', 'Mô tả', 'File FE'], socket_events)

out = os.path.abspath(OUTPUT)
wb.save(out)
print(f'Created: {out}')
print(f'Sheets: {wb.sheetnames}')
